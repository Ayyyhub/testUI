import openpyxl
import json
from datetime import datetime
from typing import List, Dict
# import requests
import requests
from utils.conf_reader import load_config
from utils.excell_reader import Excellreader


class AEUIBot:
    def __init__(self):
        self.api_url = "https://oapi.dingtalk.com/robot/send?access_token=a68d48b561a32ee60470b51e979f2dbf7b8bf4681c4fa740de9eaadb44721381"

        # self.api_url = "https://oapi.dingtalk.com/robot/send?access_token=b8f258163e1bac56cafe168c68e43fe49436126c603d90e02f3ad8247e661ecd"
        self.headers = {'Content-Type': 'application/json'}

    def format_test_results(self, format_result: List[Dict], format_sheetname: str = None) -> str:
        """格式化测试结果"""
        # 统计测试结果
        config = load_config()
        # 通过配置文件config.yaml读取到excell路径
        excell_path = config['excell_path']
        # excell_reader = Excellreader(excell_path)
        
        # 获取工作表名称列表
        workbook = openpyxl.load_workbook(excell_path)
        sheet_names = workbook.sheetnames
        # if format_sheetname and format_sheetname in sheet_names:
        #     current_sheet = workbook[format_sheetname]
        #     total_rows = current_sheet.max_row - 1  # 减去标题行

        total_cases = len(format_result)
        passed_cases = sum(1 for case in format_result if case.get('status') == 'PASS')
        failed_cases = total_cases - passed_cases
        success_rate = (passed_cases / total_cases * 100) if total_cases > 0 else 0

        # 获取失败的测试用例详情
        failed_details = [
            f"- 用例ID: {case.get('test_case_id')}\n  描述: {case.get('description')}\n  失败原因: {'请查看Log日志...'}"
            for case in format_result if case.get('status') != 'PASS'
        ]

        # 构建消息内容
        message = (
            f"执行时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            f"执行用例名称{format_sheetname}\n"
            f"- 用例总数：{total_cases}\n"
            f"- 通过用例：{passed_cases}\n"
            f"- 失败用例：{failed_cases}\n"
            f"- 成功率：{success_rate:.2f}%\n"
        )
        # 如果有失败的用例，添加失败详情
        if failed_cases > 0:
            message += "\n### 失败用例详情\n" + "\n\n".join(failed_details)

        return message

    def send_test_results(self,  result: List[Dict], sheet_name: str = None) -> bool:
        """发送测试结果到钉钉"""
        try:
            message = self.format_test_results(result,sheet_name)
            data = {
                "msgtype": "markdown",
                "markdown": {
                    "title": "AE_Cloud UI自动化测试报告",
                    "text": message
                }
            }
            
            response = requests.post(
                self.api_url,
                headers=self.headers,
                data=json.dumps(data)
            )
            
            if response.status_code == 200:
                print("测试结果发送成功")
                return True
            else:
                print(f"发送失败，状态码：{response.status_code}")
                return False
                
        except Exception as e:
            print(f"发送测试结果时发生错误：{str(e)}")
            return False

